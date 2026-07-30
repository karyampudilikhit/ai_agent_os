"""create_docx — generate a real Word document from markdown-ish content.

Same sandboxed-workspace trust tier and non-mutating rationale as
create_pptx (see that module's docstring). Understands a light
markdown subset — # / ## / ### headings, '- '/'* ' bullets, '1. '
numbered lists, plain paragraphs — so a specialist's already-written
markdown deliverable can become a real .docx without re-authoring it
in a different format.
"""

from __future__ import annotations

import re
from typing import Any, Dict

from backend.app.actions.action_registry import ActionSpec
from backend.app.actions.builtin._workspace import resolve_within, workspace_root

_NUMBERED_RE = re.compile(r"^\d+\.\s+(.*)$")


def _render_markdown_into_docx(doc, content: str) -> None:
    for raw_line in content.split("\n"):
        stripped = raw_line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith(("- ", "* ")):
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        else:
            m = _NUMBERED_RE.match(stripped)
            if m:
                doc.add_paragraph(m.group(1).strip(), style="List Number")
            else:
                # Strip stray markdown emphasis markers — python-docx
                # doesn't need them and leaving them in reads oddly.
                clean = stripped.replace("**", "").replace("__", "")
                doc.add_paragraph(clean)


def _handler(args: Dict[str, Any]) -> str:
    try:
        import docx
    except ImportError:
        return "(python-docx not installed on the server — cannot create .docx. Run: pip install python-docx)"

    filename = str(args.get("filename") or "document.docx").strip()
    if not filename.lower().endswith(".docx"):
        filename += ".docx"
    title = str(args.get("title") or "").strip()
    content = str(args.get("content") or "").strip()
    if not content:
        return "(missing 'content' — nothing to write)"

    root = workspace_root()
    try:
        target = resolve_within(root, filename)
    except ValueError as exc:
        return f"(rejected: {exc})"

    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)
    _render_markdown_into_docx(doc, content)

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    rel = target.relative_to(root)
    url = f"/api/workspace/download/{rel.as_posix()}"
    return f"Created {rel}. Download: {url}"


def _preview(args: Dict[str, Any]) -> str:
    n = len(str(args.get("content") or ""))
    label = args.get("title") or args.get("filename") or "(untitled)"
    return f"Create DOCX: {label} — {n} chars"


SPEC = ActionSpec(
    name="create_docx",
    description=(
        "Create a real Word (.docx) file in the workspace from markdown-ish text. "
        "Supports # / ## / ### headings, '- ' bullets, '1. ' numbered lists, and plain "
        "paragraphs. Writes only inside the sandboxed workspace — no approval needed. "
        "Arguments: filename (e.g. 'memo.docx'), title (optional doc title), content "
        "(the full markdown-ish body text)."
    ),
    parameters=[
        {"name": "filename", "type": "string", "description": "Output filename, e.g. 'memo.docx'", "required": True},
        {"name": "title", "type": "string", "description": "Optional document title (rendered as Title style)"},
        {"name": "content", "type": "string", "description": "Markdown-ish body text", "required": True},
    ],
    handler=_handler,
    preview=_preview,
    mutating=False,
)
